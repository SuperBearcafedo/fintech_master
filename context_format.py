# -*- coding: utf-8 -*-
import os
import sys
import traceback
from base_import import Logger
import docx
from docx import Document
from docx.shared import Length, Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_UNDERLINE
from datetime import datetime


class TableCreate:
    def __init__(self, rows, cols, file_list):
        self.document_object = Document()
        self.row = rows
        self.col = cols
        self.create_path = file_list[0]
        self.add_file_name = file_list[1]
        self.logger = Logger()

    def createTable(self):
        table = self.document_object.add_table(self.row, self.col, style="Table Grid")
        return table  # ToDo：如何不要返回到外部

    def insertText(self, table, row, col, text_content):  # 插入指定列和行
        """"
        1.行、列不可能为负数，否则打印error
        2.行、列数字小于初始化表格大小，直接插入
        3.行、列的数字>=初始化表格，-->行扩、列扩
        """
        row_count = self.row
        col_count = self.col
        if row < 0 or col < 0:
            self.logger.error("参数为负数，请重新选择行、列")

        if row < row_count and col < col_count:
            table.cell(row, col).text = text_content  # 没有超过初始设定的时候直接添加
        else:
            while row >= row_count and row != row_count - 1:
                table.add_row()  # ToDo：如何行宽一致
                row_count += 1
            while col >= col_count and col != col_count - 1:
                table.add_column(width=Cm(1))  # ToDo：如何列距一致
                col_count += 1
            table.cell(row, col).text = text_content
        self.saveTable()

        # ToDo:删除某行某列/删除整张表--> 清除某个元素 --> 修改某个元素/改某行/改某列

    def saveTable(self):
        save_path = os.path.join(self.create_path, self.add_file_name)
        self.document_object.save(save_path)  # 创建文件保存地址  ToDo：保存函数一致


if __name__ == "__main__":
    create_path = os.path.join(os.getcwd(), "project_file")
    add_file_name = "summarybook.docx"
    title_name = "第一章/First Chapter"
    chapter_name = "深度学习Python从入门到实践"

    first_table = TableCreate(1, 1, [create_path, add_file_name])
    table = first_table.createTable()
    first_table.insertText(table, 2, 2, "杂毛儿")
